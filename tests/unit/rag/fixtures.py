"""Deterministic document fixtures for the RAG tests.

Fixtures are generated at test time rather than committed as binaries, so
the test data is reproducible and reviewable. The generated documents
contain synthetic content only.
"""

from __future__ import annotations

from pathlib import Path

MARKDOWN_TEXT = """# Widget Manual

**Version:** 1.0
**Provenance:** Synthetic fixture generated for the NexusAI test suite.

## 1. Overview

The Widget is a synthetic device used to exercise document extraction.
It measures nothing of consequence and describes no real product.

## 2. Warranty

The Widget is covered by a **12-month** warranty from the date of delivery.

The warranty covers defects in materials and workmanship under normal use.

The following are **not** covered by the warranty:

- Physical damage from impact
- Liquid ingress
- Incorrect voltage supply

## 3. Calibration

Allow a 60-second warm-up period before taking calibration readings.
Calibration must be repeated every 90 days.
"""

PLAIN_TEXT = """Widget Release Notes

The Widget firmware 1.0 was released on 2026-01-15.
The release fixes an overheating issue in the sensor driver.
Support is available on weekdays between 09:00 and 17:00.
"""

DOCX_PARAGRAPHS = [
    ("Heading 1", "Widget Guide"),
    ("Normal", "The Widget is a synthetic device used in the NexusAI test suite."),
    ("Heading 2", "Safety"),
    ("Normal", "Disconnect the Widget before opening the enclosure."),
    ("Normal", "The Widget must not be operated above 40 degrees Celsius."),
    ("Heading 2", "Support"),
    ("Normal", "Contact support@example.invalid for assistance."),
]


def write_markdown(path: Path) -> Path:
    """Write the Markdown fixture."""
    path.write_text(MARKDOWN_TEXT, encoding="utf-8")
    return path


def write_text(path: Path) -> Path:
    """Write the plain-text fixture."""
    path.write_text(PLAIN_TEXT, encoding="utf-8")
    return path


def write_empty(path: Path) -> Path:
    """Write a file with no extractable text."""
    path.write_text("", encoding="utf-8")
    return path


def write_binary_garbage(path: Path) -> Path:
    """Write a file that is not a valid document of any supported type."""
    path.write_bytes(b"\x00\x01\x02not a document at all\xff\xfe")
    return path


def write_corrupt_pdf(path: Path) -> Path:
    """Write a file with a PDF extension that is not a valid PDF."""
    path.write_bytes(b"%PDF-1.7\nthis is not a real pdf body")
    return path


def write_docx(path: Path) -> Path:
    """Write a valid DOCX file using python-docx.

    Building the fixture with python-docx itself guarantees a genuine
    Office Open XML package with resolvable heading styles, rather than
    hand-rolled XML that a reader might reject.
    """
    import docx  # type: ignore[import-not-found]

    document = docx.Document()
    for style, text in DOCX_PARAGRAPHS:
        if style == "Heading 1":
            document.add_heading(text, level=1)
        elif style == "Heading 2":
            document.add_heading(text, level=2)
        else:
            document.add_paragraph(text)

    path.parent.mkdir(parents=True, exist_ok=True)
    document.save(str(path))
    return path


def write_pdf(path: Path) -> Path:
    """Write a real PDF fixture using PyMuPDF.

    The document has two pages with headings, so page numbers and heading
    extraction can both be checked.
    """
    import pymupdf  # type: ignore[import-not-found]

    document = pymupdf.open()
    pages = [
        (
            "Widget Manual\n"
            "The Widget is a synthetic device used in the NexusAI test suite.\n"
        ),
        (
            "2. Warranty\n"
            "The Widget is covered by a 12-month warranty from the date of delivery.\n"
            "The following are not covered by the warranty:\n"
            "Physical damage from impact\n"
            "Liquid ingress\n"
        ),
    ]
    for content in pages:
        page = document.new_page()
        page.insert_text((72, 72), content, fontsize=11)

    document.save(str(path))
    document.close()
    return path
